"""Turn sheet3.html into an editable A4 PowerPoint.

Background art (arch, pattern, cards, icons, table fills) is rendered to one image
with all HTML text hidden; every text block is then re-created as a native,
editable PowerPoint text box at the same position, font, size and colour.
"""
import pathlib
import re

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu

R = pathlib.Path(__file__).parent.resolve()
PX = 9525  # EMU per CSS px

COLLECT = r"""
(() => {
  const out = [];
  const done = new Set();
  const rgb = c => { const m = c.match(/[\d.]+/g); return m ? m.slice(0,3).map(Number) : [0,0,0]; };
  const style = el => {
    const cs = getComputedStyle(el);
    let col = rgb(cs.color);
    if (cs.color.includes('rgba(0, 0, 0, 0)') || el.closest('.foil')) col = [214,172,94];
    return {font: cs.fontFamily.split(',')[0].replace(/["']/g,''), size: parseFloat(cs.fontSize),
            weight: cs.fontWeight, italic: cs.fontStyle === 'italic', color: col,
            ls: cs.letterSpacing === 'normal' ? 0 : parseFloat(cs.letterSpacing),
            upper: cs.textTransform === 'uppercase'};
  };
  const hasOwnText = el => [...el.childNodes].some(n => n.nodeType === 3 && n.textContent.trim());
  document.querySelectorAll('body *').forEach(el => {
    if (el.closest('svg') || done.has(el)) return;
    if ([...done].some(d => d.contains(el))) return;
    if (!hasOwnText(el)) return;
    // measure the text itself (not icons / flex boxes)
    let L=1e9,T=1e9,Rr=-1e9,B=-1e9;
    el.childNodes.forEach(n => {
      if (n.nodeName === 'svg' || n.nodeName === 'BR') return;
      if (n.nodeType === 1 && n.querySelector && n.querySelector('svg')) return;
      if (!n.textContent.trim()) return;
      const rg = document.createRange(); rg.selectNodeContents(n.nodeType===3 ? n.parentNode : n);
      if (n.nodeType === 3) { rg.setStart(n,0); rg.setEnd(n,n.length); }
      for (const q of rg.getClientRects()) { if (q.width<0.5) continue; L=Math.min(L,q.left);T=Math.min(T,q.top);Rr=Math.max(Rr,q.right);B=Math.max(B,q.bottom); }
    });
    const r = {left:L, top:T, width:Rr-L, height:B-T};
    if (!(r.width > 0)) return;
    const cs = getComputedStyle(el);
    const lines = [[]];
    el.childNodes.forEach(n => {
      if (n.nodeType === 3) { const t = n.textContent.replace(/\s+/g, ' '); if (t.trim() || t === ' ') lines[lines.length-1].push({text: t, ...style(el)}); }
      else if (n.nodeName === 'BR') lines.push([]);
      else if (n.nodeType === 1 && n.textContent.trim()) lines[lines.length-1].push({text: n.textContent.replace(/\s+/g,' '), ...style(n)});
    });
    let align = cs.textAlign;
    if (cs.display.includes('flex') && cs.justifyContent === 'center') align = 'center';
    const single = lines.length === 1 && r.height < parseFloat(cs.fontSize) * 1.9;
    out.push({x: r.left, y: r.top, w: r.width, h: r.height, align, single, lh: parseFloat(cs.lineHeight) || null, lines});
    done.add(el);
  });
  // hide all collected text for the background render
  document.querySelectorAll('.foil').forEach(f => { f.style.background = 'none'; });
  done.forEach(el => { el.style.color = 'transparent'; el.style.webkitTextFillColor = 'transparent';
    el.querySelectorAll('*').forEach(c => { if (!c.closest('svg')) { c.style.color = 'transparent'; c.style.webkitTextFillColor = 'transparent'; } }); });
  return out;
})()
"""

FONT_MAP = {"Cinzel": "Cinzel", "Serif": "Cormorant Garamond", "SerifI": "Cormorant Garamond", "Sans": "Montserrat"}


def main():
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
        pg = b.new_page(viewport={"width": 794, "height": 1123}, device_scale_factor=3)
        pg.goto((R / "sheet3.html").as_uri())
        pg.evaluate("document.fonts.ready.then(()=>true)")
        pg.wait_for_timeout(400)
        blocks = pg.evaluate(COLLECT)
        pg.wait_for_timeout(200)
        pg.screenshot(path=str(R / "pptx-background.png"))
        b.close()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(794 * PX), Emu(1123 * PX)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(R / "pptx-background.png"), 0, 0, prs.slide_width, prs.slide_height)

    for bl in blocks:
        pad = 6 if bl["single"] else max(6, bl["w"] * 0.06)
        tb = s.shapes.add_textbox(Emu(int((bl["x"] - pad) * PX)), Emu(int(bl["y"] * PX)),
                                  Emu(int((bl["w"] + 2 * pad) * PX)), Emu(int(bl["h"] * PX)))
        tf = tb.text_frame
        tf.word_wrap = (not bl["single"]) and len(bl["lines"]) == 1
        tf.margin_left = tf.margin_right = Emu(pad * PX)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if bl["single"] else MSO_ANCHOR.TOP
        for li, runs in enumerate(bl["lines"]):
            para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER if bl["single"] else {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(bl["align"], PP_ALIGN.LEFT)
            if bl["lh"] and not bl["single"]:
                para.line_spacing = Emu(int(bl["lh"] * PX))
            for rn in runs:
                t = rn["text"]
                if li == 0 and rn is runs[0]:
                    t = t.lstrip()
                if rn is runs[-1]:
                    t = t.rstrip()
                if rn["upper"]:
                    t = t.upper()
                if not t:
                    continue
                r = para.add_run()
                r.text = t
                f = r.font
                fam = rn["font"]
                f.name = FONT_MAP.get(fam, fam)
                f.size = Emu(int(rn["size"] * 0.75 * 12700))
                f.bold = str(rn["weight"]) in ("600", "700", "800", "bold")
                f.italic = rn["italic"] or fam == "SerifI"
                f.color.rgb = RGBColor(*[int(c) for c in rn["color"]])
                if rn["ls"]:
                    r._r.get_or_add_rPr().set("spc", str(int(rn["ls"] * 0.75 * 100)))
    prs.save(str(R / "PL NILAYA (editable).pptx"))
    print(len(blocks), "text boxes")


main()
